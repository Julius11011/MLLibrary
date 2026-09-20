# generate_phase1_complete.py
import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from core_builder import (
    create_docx_builder, add_header_box, add_heading_1, add_heading_2,
    add_body_p, add_student_explanation_box, add_alac_box,
    add_latin_maxims_box, add_statcon_table
)

from statcon_database_phase1 import cases_p1
from generate_phase1_full import cases_p1_part2
from statcon_database_phase1_complete import cases_p1_part3

# Full cases list
all_p1_cases = []
seen_nums = set()
for c in cases_p1 + cases_p1_part2 + cases_p1_part3:
    if c['num'] not in seen_nums:
        all_p1_cases.append(c)
        seen_nums.add(c['num'])

all_p1_cases.sort(key=lambda x: x['num'])
print(f"Total Unique Phase 1 Cases: {len(all_p1_cases)}")

def build_phase1_documents():
    txt_output_path = r"C:\Users\JR\Downloads\14All-All41\MLC\First Sem 1st Year\Statutory Construction\StatCons\StatCons_Phase1_Chapters_I_II.txt"
    docx_output_path = r"C:\Users\JR\Downloads\14All-All41\MLC\First Sem 1st Year\Statutory Construction\StatCons\StatCons_Phase1_Chapters_I_II.docx"
    
    # 1. Build Text File
    with open(txt_output_path, "w", encoding="utf-8") as f:
        f.write("MANILA LAW COLLEGE\n")
        f.write("Subject: Statutory Construction (2 units)\n")
        f.write("Instructor: Atty. Jason P. Raposas\n")
        f.write("Reference Syllabus: StatCon_Syllabus_JPR.pdf\n\n")
        f.write("================================================================================\n")
        f.write("PHASE 1: CHAPTERS I & II CASE DIGESTS IN ALAC FORMAT & STATCON ANALYSIS\n")
        f.write("(With Latin Maxims + Inline English Translations + Recitation Guides)\n")
        f.write("================================================================================\n\n")
        
        for c in all_p1_cases:
            f.write("="*80 + "\n")
            f.write(f"{c['num']}. {c['title'].upper()}\n")
            f.write(f"{c['citation']}\n")
            f.write("="*80 + "\n\n")
            
            f.write("I. PARTIES & PROCEDURAL BACKGROUND\n")
            f.write(f"* Plaintiff/Petitioner: {c['plaintiff']}\n")
            f.write(f"* Defendant/Respondent: {c['defendant']}\n")
            f.write(f"* Originating Court: {c['court']}\n")
            f.write(f"* Nature of the Action: {c['nature']}\n\n")
            
            f.write("II. WHOLE FACTS\n")
            f.write(f"{c['facts']}\n\n")
            f.write(f"Plain Meaning Summary of Facts:\n{c['plain_facts']}\n\n")
            
            f.write("-" * 80 + "\n")
            f.write("III. THE LEGAL ISSUE\n")
            f.write("-" * 80 + "\n")
            f.write(f"{c['issue']}\n\n")
            f.write("[HOW TO TACKLE IN RECITATION / SIMPLE WORDING: THE LEGAL ISSUE]\n")
            f.write(f"* Plain Meaning Explanation of the Issue:\n  {c['issue_plain']}\n")
            f.write(f"* How to Tackle in Recitation When Called:\n  {c['issue_recite']}\n\n")
            
            f.write("-" * 80 + "\n")
            f.write("IV. ALAC DIGEST\n")
            f.write("-" * 80 + "\n\n")
            
            f.write("[A] ANSWER:\n")
            f.write(f"{c['answer']}\n\n")
            f.write("[HOW TO TACKLE IN RECITATION / SIMPLE WORDING: THE ANSWER]\n")
            f.write(f"* Plain Meaning Explanation of the Answer:\n  {c['answer_plain']}\n")
            f.write(f"* How to Tackle in Recitation When Called:\n  {c['answer_recite']}\n\n")
            
            f.write("-" * 80 + "\n")
            f.write("[L] LEGAL BASIS:\n")
            for lbl, val in c['legal_basis']:
                f.write(f"{lbl} {val}\n")
            f.write("\n[HOW TO TACKLE IN RECITATION / SIMPLE WORDING: LEGAL BASIS]\n")
            f.write(f"* Plain Meaning Explanation of the Legal Basis:\n  {c['legal_basis_plain']}\n")
            f.write(f"* How to Tackle in Recitation When Called:\n  {c['legal_basis_recite']}\n\n")
            
            f.write("-" * 80 + "\n")
            f.write("[A] ANALYSIS / APPLICATION:\n")
            for lbl, val in c['analysis']:
                f.write(f"{lbl} {val}\n")
            f.write("\n[HOW TO TACKLE IN RECITATION / SIMPLE WORDING: ANALYSIS / APPLICATION]\n")
            f.write(f"* Plain Meaning Explanation of the Analysis:\n  {c['analysis_plain']}\n")
            f.write(f"* How to Tackle in Recitation When Called:\n  {c['analysis_recite']}\n\n")
            
            f.write("-" * 80 + "\n")
            f.write("[C] CONCLUSION:\n")
            f.write(f"{c['conclusion']}\n\n")
            f.write("[HOW TO TACKLE IN RECITATION / SIMPLE WORDING: CONCLUSION]\n")
            f.write(f"* Plain Meaning Explanation of the Conclusion:\n  {c['conclusion_plain']}\n")
            f.write(f"* How to Tackle in Recitation When Called:\n  {c['conclusion_recite']}\n\n")
            
            f.write("-" * 80 + "\n")
            f.write("V. RELATION TO STATUTORY CONSTRUCTION (StatCon_Syllabus_JPR.pdf)\n")
            f.write("-" * 80 + "\n\n")
            f.write(f"Applicable Syllabus Topics:\n{c['syllabus']}\n\n")
            
            f.write("LATIN MAXIMS WITH ENGLISH TRANSLATION & LEGAL MEANING:\n")
            for idx, (m_latin, m_trans, m_mean) in enumerate(c['latin_maxims'], 1):
                f.write(f"{idx}. {m_latin}\n")
                f.write(f"   • Literal English Translation: {m_trans}\n")
                f.write(f"   • Plain Legal Meaning & StatCon Application: {m_mean}\n")
            f.write("\n")
            
            f.write(f"Case Violation in StatCon:\n{c['violation_text']}\n\n")
            f.write(f"Doctrinal Execution / Principle:\n{c['execution_text']}\n\n")
            
            f.write("[HOW TO TACKLE IN RECITATION / SIMPLE WORDING: RELATION TO STATUTORY CONSTRUCTION]\n")
            f.write(f"* Plain Meaning Explanation of the StatCon Relation:\n  {c['statcon_plain']}\n")
            f.write(f"* How to Tackle in Recitation When Called:\n  {c['statcon_recite']}\n\n\n")
    
    print("Successfully generated Phase 1 TXT file:", txt_output_path)
    
    # 2. Build Word File
    doc = create_docx_builder()
    add_header_box(
        doc,
        title="PHASE 1: CHAPTERS I & II CASE DIGESTS COMPILATION",
        subtitle="Full ALAC Digests with Latin Maxims (Inline English) & Recitation Guides",
        course_info="MANILA LAW COLLEGE | Statutory Construction (2 units) | Atty. Jason P. Raposas"
    )
    
    for idx, c in enumerate(all_p1_cases):
        if idx > 0:
            doc.add_page_break()
            
        add_heading_1(doc, f"{c['num']}. {c['title']}")
        add_body_p(doc, c['citation'], bold_prefix="Citation: ", italic=True)
        
        add_heading_2(doc, "I. Parties & Procedural Background")
        add_body_p(doc, c['plaintiff'], bold_prefix="• Plaintiff/Petitioner: ")
        add_body_p(doc, c['defendant'], bold_prefix="• Defendant/Respondent: ")
        add_body_p(doc, c['court'], bold_prefix="• Originating Court: ")
        add_body_p(doc, c['nature'], bold_prefix="• Nature of the Action: ")
        
        add_heading_2(doc, "II. Whole Facts")
        add_body_p(doc, c['facts'])
        add_body_p(doc, c['plain_facts'], bold_prefix="Plain Meaning Summary: ", italic=True)
        
        add_heading_2(doc, "III. The Legal Issue")
        add_body_p(doc, c['issue'])
        add_student_explanation_box(
            doc,
            "The Legal Issue",
            plain_explanation=c['issue_plain'],
            recitation_guide=c['issue_recite']
        )
        
        add_heading_2(doc, "IV. ALAC Digest")
        add_alac_box(doc, "A", "Answer", [c['answer']])
        add_student_explanation_box(
            doc,
            "The Answer",
            plain_explanation=c['answer_plain'],
            recitation_guide=c['answer_recite']
        )
        
        add_alac_box(doc, "L", "Legal Basis", c['legal_basis'])
        add_student_explanation_box(
            doc,
            "Legal Basis",
            plain_explanation=c['legal_basis_plain'],
            recitation_guide=c['legal_basis_recite']
        )
        
        add_alac_box(doc, "A", "Analysis / Application", c['analysis'])
        add_student_explanation_box(
            doc,
            "Analysis / Application",
            plain_explanation=c['analysis_plain'],
            recitation_guide=c['analysis_recite']
        )
        
        add_alac_box(doc, "C", "Conclusion", [c['conclusion']])
        add_student_explanation_box(
            doc,
            "Conclusion",
            plain_explanation=c['conclusion_plain'],
            recitation_guide=c['conclusion_recite']
        )
        
        add_heading_2(doc, "V. Relation to Statutory Construction")
        add_latin_maxims_box(doc, c['latin_maxims'])
        add_statcon_table(
            doc,
            syllabus_topic=c['syllabus'],
            primary_maxim=c['primary_maxim'],
            secondary_maxims=c['secondary_maxims'],
            violation_text=c['violation_text'],
            execution_text=c['execution_text']
        )
        add_student_explanation_box(
            doc,
            "Relation to Statutory Construction",
            plain_explanation=c['statcon_plain'],
            recitation_guide=c['statcon_recite']
        )
    
    doc.save(docx_output_path)
    print("Successfully generated Phase 1 DOCX file:", docx_output_path)

if __name__ == "__main__":
    build_phase1_documents()
