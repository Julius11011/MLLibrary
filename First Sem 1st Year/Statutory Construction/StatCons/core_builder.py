# core_builder.py
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_docx_builder():
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
    return doc

# Color Palette
COLOR_PRIMARY = RGBColor(15, 32, 67)      # Deep Navy
COLOR_SECONDARY = RGBColor(120, 30, 30)  # Deep Crimson
COLOR_STUDENT = RGBColor(13, 110, 75)    # Forest Emerald Green
COLOR_LATIN = RGBColor(140, 50, 0)       # Amber/Brown
COLOR_FORMAT_HDR = RGBColor(40, 55, 90)  # Slate Navy for Template
COLOR_TEXT = RGBColor(33, 37, 41)        # Dark Charcoal

HEX_PRIMARY = "0F2043"
HEX_SECONDARY = "781E1E"
HEX_LIGHT_BG = "F4F6F9"
HEX_STUDENT_BG = "EBF7F0"
HEX_LATIN_BG = "FDF6EE"
HEX_TEMPLATE_BG = "F0F4FA"
HEX_CALLOUT_BG = "F8F9FB"
HEX_VERDICT_BG = "EEF6F0"
HEX_DOCTRINE_BG = "FBF6EE"

def add_header_box(doc, title, subtitle, course_info):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, HEX_PRIMARY)
    set_cell_margins(cell, top=200, bottom=200, left=240, right=240)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    run1 = p.add_run("MANILA LAW COLLEGE\n")
    run1.font.name = "Arial"
    run1.font.size = Pt(14)
    run1.font.bold = True
    run1.font.color.rgb = RGBColor(255, 255, 255)
    
    run2 = p.add_run(f"{course_info}\n")
    run2.font.name = "Arial"
    run2.font.size = Pt(10)
    run2.font.color.rgb = RGBColor(210, 225, 245)
    
    run3 = p.add_run(f"{title}\n")
    run3.font.name = "Arial"
    run3.font.size = Pt(12)
    run3.font.bold = True
    run3.font.color.rgb = RGBColor(255, 215, 0)
    
    run4 = p.add_run(subtitle)
    run4.font.name = "Arial"
    run4.font.size = Pt(9.5)
    run4.font.italic = True
    run4.font.color.rgb = RGBColor(230, 235, 245)

def add_chapter_banner(doc, chapter_num, chapter_title, chapter_subtitle, syllabus_topics=None):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, HEX_PRIMARY)
    set_cell_margins(cell, top=160, bottom=160, left=200, right=200)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    
    r_ch = p.add_run(f"CHAPTER {chapter_num}\n")
    r_ch.font.name = "Arial"
    r_ch.font.size = Pt(13)
    r_ch.font.bold = True
    r_ch.font.color.rgb = RGBColor(255, 215, 0)
    
    r_title = p.add_run(f"{chapter_title.upper()}\n")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(12)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(255, 255, 255)
    
    r_sub = p.add_run(chapter_subtitle)
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(9.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(215, 230, 250)
    
    if syllabus_topics:
        p_top = cell.add_paragraph()
        p_top.paragraph_format.space_before = Pt(6)
        p_top.paragraph_format.space_after = Pt(2)
        p_top.paragraph_format.line_spacing = 1.15
        r_top_lbl = p_top.add_run("Key Syllabus Coverage:\n")
        r_top_lbl.font.name = "Arial"
        r_top_lbl.font.size = Pt(9.5)
        r_top_lbl.font.bold = True
        r_top_lbl.font.color.rgb = RGBColor(255, 215, 0)
        
        for topic in syllabus_topics:
            p_t = cell.add_paragraph()
            p_t.paragraph_format.space_before = Pt(1)
            p_t.paragraph_format.space_after = Pt(1)
            r_t = p_t.add_run(f"• {topic}")
            r_t.font.name = "Calibri"
            r_t.font.size = Pt(9)
            r_t.font.color.rgb = RGBColor(240, 245, 255)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = COLOR_SECONDARY
    return p

def add_body_p(doc, text, bold_prefix="", italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_TEXT
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)
    r.font.italic = italic
    r.font.color.rgb = COLOR_TEXT
    return p

def add_student_explanation_box(doc, section_title, plain_explanation, recitation_guide):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, HEX_STUDENT_BG)
    set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    r_head = p.add_run(f"💡 HOW TO TACKLE IN RECITATION / SIMPLE WORDING: {section_title.upper()}\n")
    r_head.font.name = "Arial"
    r_head.font.size = Pt(9.5)
    r_head.font.bold = True
    r_head.font.color.rgb = COLOR_STUDENT
    
    p1 = cell.add_paragraph()
    p1.paragraph_format.space_before = Pt(1)
    p1.paragraph_format.space_after = Pt(2)
    p1.paragraph_format.line_spacing = 1.15
    r_p1_lbl = p1.add_run("• Plain Meaning Explanation: ")
    r_p1_lbl.font.name = "Calibri"
    r_p1_lbl.font.size = Pt(9.5)
    r_p1_lbl.font.bold = True
    r_p1_lbl.font.color.rgb = COLOR_STUDENT
    r_p1_val = p1.add_run(plain_explanation)
    r_p1_val.font.name = "Calibri"
    r_p1_val.font.size = Pt(9.5)
    r_p1_val.font.color.rgb = COLOR_TEXT
    
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(1)
    p2.paragraph_format.space_after = Pt(2)
    p2.paragraph_format.line_spacing = 1.15
    r_p2_lbl = p2.add_run("• How to Tackle in Recitation: ")
    r_p2_lbl.font.name = "Calibri"
    r_p2_lbl.font.size = Pt(9.5)
    r_p2_lbl.font.bold = True
    r_p2_lbl.font.color.rgb = COLOR_STUDENT
    r_p2_val = p2.add_run(recitation_guide)
    r_p2_val.font.name = "Calibri"
    r_p2_val.font.size = Pt(9.5)
    r_p2_val.font.italic = True
    r_p2_val.font.color.rgb = COLOR_TEXT

def add_alac_box(doc, letter, letter_title, content_list):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, HEX_LIGHT_BG)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    r_letter = p.add_run(f"[{letter}] {letter_title.upper()}\n")
    r_letter.font.name = "Arial"
    r_letter.font.size = Pt(10.5)
    r_letter.font.bold = True
    r_letter.font.color.rgb = COLOR_PRIMARY
    
    for item in content_list:
        p_item = cell.add_paragraph()
        p_item.paragraph_format.space_before = Pt(1)
        p_item.paragraph_format.space_after = Pt(2)
        p_item.paragraph_format.line_spacing = 1.15
        if isinstance(item, tuple):
            r_head = p_item.add_run(item[0] + " ")
            r_head.font.name = "Calibri"
            r_head.font.size = Pt(10)
            r_head.font.bold = True
            r_head.font.color.rgb = COLOR_TEXT
            
            r_body = p_item.add_run(item[1])
            r_body.font.name = "Calibri"
            r_body.font.size = Pt(10)
            r_body.font.color.rgb = COLOR_TEXT
        else:
            r_body = p_item.add_run(str(item))
            r_body.font.name = "Calibri"
            r_body.font.size = Pt(10)
            r_body.font.color.rgb = COLOR_TEXT

def add_latin_maxims_box(doc, maxims_list):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, HEX_LATIN_BG)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    r_head = p.add_run("📜 GOVERNING LATIN MAXIMS WITH ENGLISH TRANSLATIONS\n")
    r_head.font.name = "Arial"
    r_head.font.size = Pt(10)
    r_head.font.bold = True
    r_head.font.color.rgb = COLOR_LATIN
    
    for idx, (latin_with_eng, translation, meaning) in enumerate(maxims_list, 1):
        p_item = cell.add_paragraph()
        p_item.paragraph_format.space_before = Pt(2)
        p_item.paragraph_format.space_after = Pt(3)
        p_item.paragraph_format.line_spacing = 1.15
        
        r_num = p_item.add_run(f"{idx}. {latin_with_eng}\n")
        r_num.font.name = "Calibri"
        r_num.font.size = Pt(10)
        r_num.font.bold = True
        r_num.font.color.rgb = COLOR_LATIN
        
        r_tr_lbl = p_item.add_run("   • Literal English Translation: ")
        r_tr_lbl.font.name = "Calibri"
        r_tr_lbl.font.size = Pt(9.5)
        r_tr_lbl.font.bold = True
        r_tr_lbl.font.color.rgb = COLOR_TEXT
        
        r_tr_val = p_item.add_run(f"{translation}\n")
        r_tr_val.font.name = "Calibri"
        r_tr_val.font.size = Pt(9.5)
        r_tr_val.font.italic = True
        r_tr_val.font.color.rgb = COLOR_TEXT
        
        r_mn_lbl = p_item.add_run("   • Plain Legal Meaning & StatCon Application: ")
        r_mn_lbl.font.name = "Calibri"
        r_mn_lbl.font.size = Pt(9.5)
        r_mn_lbl.font.bold = True
        r_mn_lbl.font.color.rgb = COLOR_TEXT
        
        r_mn_val = p_item.add_run(f"{meaning}")
        r_mn_val.font.name = "Calibri"
        r_mn_val.font.size = Pt(9.5)
        r_mn_val.font.color.rgb = COLOR_TEXT

def add_statcon_table(doc, syllabus_topic, primary_maxim, secondary_maxims, violation_text, execution_text):
    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    data = [
        ("Syllabus Topic Alignment\n(StatCon_Syllabus_JPR.pdf)", syllabus_topic),
        ("Governing Latin Maxims & English", f"Primary:\n{primary_maxim}\n\nSecondary:\n{secondary_maxims}"),
        ("Case Violation in StatCon", violation_text),
        ("Doctrinal Execution / Principle", execution_text)
    ]
    
    for i, (label, val) in enumerate(data):
        cell_lbl = table.cell(i, 0)
        cell_val = table.cell(i, 1)
        
        set_cell_background(cell_lbl, "EAEFF5")
        set_cell_background(cell_val, "FAFAFA")
        set_cell_margins(cell_lbl, top=80, bottom=80, left=120, right=120)
        set_cell_margins(cell_val, top=80, bottom=80, left=120, right=120)
        
        p_lbl = cell_lbl.paragraphs[0]
        p_lbl.paragraph_format.space_before = Pt(2)
        p_lbl.paragraph_format.space_after = Pt(2)
        r_lbl = p_lbl.add_run(label)
        r_lbl.font.name = "Arial"
        r_lbl.font.size = Pt(9)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = COLOR_PRIMARY
        
        p_val = cell_val.paragraphs[0]
        p_val.paragraph_format.space_before = Pt(2)
        p_val.paragraph_format.space_after = Pt(2)
        p_val.paragraph_format.line_spacing = 1.15
        r_val = p_val.add_run(val)
        r_val.font.name = "Calibri"
        r_val.font.size = Pt(9.5)
        r_val.font.color.rgb = COLOR_TEXT

# ==============================================================================
# LAYER B: STANDARDIZED 12-SECTION STATUTORY CONSTRUCTION CASE DIGEST FORMAT
# ==============================================================================

def add_template_banner(doc):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "1E3A5F") # Dark Navy Blue
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    
    r_hdr = p.add_run("⚖ STANDARDIZED CASE DIGEST: STATUTORY CONSTRUCTION FORMAT\n")
    r_hdr.font.name = "Arial"
    r_hdr.font.size = Pt(11)
    r_hdr.font.bold = True
    r_hdr.font.color.rgb = RGBColor(255, 215, 0)
    
    r_sub = p.add_run("(Complete 12-Section Academic & Judicial Construction Analysis)")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(9)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(220, 235, 255)

def add_template_meta_table(doc, title, gr_no, date, ponente):
    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    rows_data = [
        ("Case Title:", title),
        ("G.R. No. / Record:", gr_no),
        ("Date of Promulgation:", date),
        ("Ponente / Court:", ponente)
    ]
    
    for i, (lbl, val) in enumerate(rows_data):
        c_lbl = table.cell(i, 0)
        c_val = table.cell(i, 1)
        
        set_cell_background(c_lbl, "EBF0F7")
        set_cell_background(c_val, "FAFCFF")
        set_cell_margins(c_lbl, top=60, bottom=60, left=120, right=120)
        set_cell_margins(c_val, top=60, bottom=60, left=120, right=120)
        
        p0 = c_lbl.paragraphs[0]
        p0.paragraph_format.space_before = Pt(1)
        p0.paragraph_format.space_after = Pt(1)
        r0 = p0.add_run(lbl)
        r0.font.name = "Arial"
        r0.font.size = Pt(9.5)
        r0.font.bold = True
        r0.font.color.rgb = COLOR_PRIMARY
        
        p1 = c_val.paragraphs[0]
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        r1 = p1.add_run(val)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.bold = (i == 0)
        r1.font.color.rgb = COLOR_TEXT

def add_template_section_h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY
    return p

def add_template_section_h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = COLOR_SECONDARY
    return p

def add_template_box(doc, title, lines_list, hex_bg=HEX_CALLOUT_BG, title_color=COLOR_PRIMARY):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, hex_bg)
    set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    if title:
        r_head = p.add_run(f"{title}\n")
        r_head.font.name = "Arial"
        r_head.font.size = Pt(9.5)
        r_head.font.bold = True
        r_head.font.color.rgb = title_color
        
    for item in lines_list:
        p_item = cell.add_paragraph()
        p_item.paragraph_format.space_before = Pt(1)
        p_item.paragraph_format.space_after = Pt(2)
        p_item.paragraph_format.line_spacing = 1.15
        if isinstance(item, tuple):
            r_lbl = p_item.add_run(item[0] + " ")
            r_lbl.font.name = "Calibri"
            r_lbl.font.size = Pt(9.5)
            r_lbl.font.bold = True
            r_lbl.font.color.rgb = COLOR_TEXT
            
            r_val = p_item.add_run(str(item[1]))
            r_val.font.name = "Calibri"
            r_val.font.size = Pt(9.5)
            r_val.font.color.rgb = COLOR_TEXT
        else:
            r_val = p_item.add_run(str(item))
            r_val.font.name = "Calibri"
            r_val.font.size = Pt(9.5)
            r_val.font.color.rgb = COLOR_TEXT

def add_template_concepts_table(doc, concepts_list):
    table = doc.add_table(rows=len(concepts_list), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    for i, c in enumerate(concepts_list):
        cell_lbl = table.cell(i, 0)
        cell_val = table.cell(i, 1)
        
        bg_col = "F5F8FC" if i % 2 == 0 else "FFFFFF"
        set_cell_background(cell_lbl, "EBF0F8")
        set_cell_background(cell_val, bg_col)
        set_cell_margins(cell_lbl, top=70, bottom=70, left=100, right=100)
        set_cell_margins(cell_val, top=70, bottom=70, left=120, right=120)
        
        p_lbl = cell_lbl.paragraphs[0]
        p_lbl.paragraph_format.space_before = Pt(1)
        p_lbl.paragraph_format.space_after = Pt(1)
        r_lbl = p_lbl.add_run(f"{c['letter']}. {c['name']}")
        r_lbl.font.name = "Arial"
        r_lbl.font.size = Pt(9)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = COLOR_PRIMARY
        
        p_val = cell_val.paragraphs[0]
        p_val.paragraph_format.space_before = Pt(1)
        p_val.paragraph_format.space_after = Pt(1)
        p_val.paragraph_format.line_spacing = 1.15
        
        app_tag = " [APPLICABLE] " if c.get('applicable', True) else " [NOT APPLICABLE] "
        r_tag = p_val.add_run(app_tag)
        r_tag.font.name = "Calibri"
        r_tag.font.size = Pt(8.5)
        r_tag.font.bold = True
        r_tag.font.color.rgb = COLOR_STUDENT if c.get('applicable', True) else COLOR_SECONDARY
        
        r_exp = p_val.add_run(c.get('explanation', ''))
        r_exp.font.name = "Calibri"
        r_exp.font.size = Pt(9)
        r_exp.font.color.rgb = COLOR_TEXT

def add_template_verdict_box(doc, verdict_type, explanation_text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, HEX_VERDICT_BG)
    set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    r_hdr = p.add_run("X. STATUTORY CONSTRUCTION OR JUDICIAL LEGISLATION?\n")
    r_hdr.font.name = "Arial"
    r_hdr.font.size = Pt(9.5)
    r_hdr.font.bold = True
    r_hdr.font.color.rgb = COLOR_STUDENT
    
    p_chk = cell.add_paragraph()
    p_chk.paragraph_format.space_before = Pt(1)
    p_chk.paragraph_format.space_after = Pt(3)
    
    c1 = "☑" if "proper" in verdict_type.lower() else "☐"
    c2 = "☑" if "judicial legislation" in verdict_type.lower() and "proper" not in verdict_type.lower() else "☐"
    c3 = "☑" if "debatable" in verdict_type.lower() else "☐"
    
    r_chk = p_chk.add_run(f"{c1} Proper statutory construction    {c2} Judicial legislation    {c3} Debatable")
    r_chk.font.name = "Arial"
    r_chk.font.size = Pt(9.5)
    r_chk.font.bold = True
    r_chk.font.color.rgb = COLOR_PRIMARY
    
    p_exp = cell.add_paragraph()
    p_exp.paragraph_format.space_before = Pt(2)
    p_exp.paragraph_format.space_after = Pt(2)
    p_exp.paragraph_format.line_spacing = 1.15
    r_exp_lbl = p_exp.add_run("Explanation (3–5 sentences):\n")
    r_exp_lbl.font.name = "Calibri"
    r_exp_lbl.font.size = Pt(9.5)
    r_exp_lbl.font.bold = True
    r_exp_lbl.font.color.rgb = COLOR_TEXT
    
    r_exp_val = p_exp.add_run(explanation_text)
    r_exp_val.font.name = "Calibri"
    r_exp_val.font.size = Pt(9.5)
    r_exp_val.font.color.rgb = COLOR_TEXT

def add_template_doctrine_takeaway(doc, doctrine_text, takeaway_text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, HEX_DOCTRINE_BG)
    set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    r_doc_lbl = p.add_run("XI. DOCTRINE (Statutory Construction Principle):\n")
    r_doc_lbl.font.name = "Arial"
    r_doc_lbl.font.size = Pt(9.5)
    r_doc_lbl.font.bold = True
    r_doc_lbl.font.color.rgb = COLOR_LATIN
    
    r_doc_val = p.add_run(f"{doctrine_text}\n\n")
    r_doc_val.font.name = "Calibri"
    r_doc_val.font.size = Pt(9.5)
    r_doc_val.font.color.rgb = COLOR_TEXT
    
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(1)
    p2.paragraph_format.space_after = Pt(2)
    r_tk_lbl = p2.add_run("XII. ONE-SENTENCE TAKEAWAY:\n")
    r_tk_lbl.font.name = "Arial"
    r_tk_lbl.font.size = Pt(9.5)
    r_tk_lbl.font.bold = True
    r_tk_lbl.font.color.rgb = COLOR_LATIN
    
    r_tk_val = p2.add_run(takeaway_text)
    r_tk_val.font.name = "Calibri"
    r_tk_val.font.size = Pt(9.5)
    r_tk_val.font.italic = True
    r_tk_val.font.color.rgb = COLOR_TEXT

print("Enhanced core builder with complete 12-section template elements ready.")
