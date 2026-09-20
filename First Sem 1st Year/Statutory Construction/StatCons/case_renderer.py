# case_renderer.py
"""
Unified Case Renderer for Statutory Construction Compendium.
Renders any case dictionary into both Layer A (ALAC & Recitation Suite)
and Layer B (12-Section Case Digest Format) for DOCX and TXT.
"""

from docx.shared import Pt, Inches, RGBColor
from core_builder import (
    add_heading_1, add_heading_2, add_body_p,
    add_student_explanation_box, add_alac_box,
    add_latin_maxims_box, add_statcon_table,
    add_template_banner, add_template_meta_table,
    add_template_section_h1, add_template_section_h2,
    add_template_box, add_template_concepts_table,
    add_template_verdict_box, add_template_doctrine_takeaway,
    COLOR_PRIMARY, COLOR_SECONDARY, COLOR_STUDENT, COLOR_LATIN, COLOR_TEXT
)
from statcon_format_engine import generate_case_digest_format_data

def render_case_to_docx(doc, case):
    """Renders a single case to python-docx Document with both layers."""
    tmpl = generate_case_digest_format_data(case)
    
    # Case Header
    add_heading_1(doc, f"{case['num']}. {case['title'].upper()}")
    add_body_p(doc, case['citation'], bold_prefix="Citation: ", italic=True)
    
    # =========================================================================
    # LAYER A: ALAC DIGEST & RECITATION SUITE
    # =========================================================================
    add_heading_2(doc, "LAYER A: ALAC DIGEST, RECITATION SCRIPT & STATCON ANALYSIS")
    
    # Section I
    add_heading_2(doc, "I. PARTIES & PROCEDURAL BACKGROUND")
    add_body_p(doc, case.get('plaintiff', 'N/A'), bold_prefix="• Plaintiff/Petitioner: ")
    add_body_p(doc, case.get('defendant', 'N/A'), bold_prefix="• Defendant/Respondent: ")
    add_body_p(doc, case.get('court', 'N/A'), bold_prefix="• Originating Court: ")
    add_body_p(doc, case.get('nature', 'N/A'), bold_prefix="• Nature of the Action: ")
    
    # Section II
    add_heading_2(doc, "II. WHOLE FACTS")
    add_body_p(doc, case.get('facts', ''))
    add_student_explanation_box(
        doc,
        "Summary of Facts",
        case.get('plain_facts', ''),
        "When reciting the facts, state the core dispute in 2 sentences, then state how the case reached the Supreme Court."
    )
    
    # Section III
    add_heading_2(doc, "III. THE LEGAL ISSUE")
    add_body_p(doc, case.get('issue', ''))
    add_student_explanation_box(
        doc,
        "The Legal Issue",
        case.get('issue_plain', ''),
        case.get('issue_recite', '')
    )
    
    # Section IV
    add_heading_2(doc, "IV. ALAC DIGEST")
    
    # [A] Answer
    add_alac_box(doc, "A", "Answer", [case.get('answer', '')])
    add_student_explanation_box(
        doc,
        "The Answer",
        case.get('answer_plain', ''),
        case.get('answer_recite', '')
    )
    
    # [L] Legal Basis
    add_alac_box(doc, "L", "Legal Basis", case.get('legal_basis', []))
    add_student_explanation_box(
        doc,
        "Legal Basis",
        case.get('legal_basis_plain', ''),
        case.get('legal_basis_recite', '')
    )
    
    # [A] Analysis
    add_alac_box(doc, "A", "Analysis / Application", case.get('analysis', []))
    add_student_explanation_box(
        doc,
        "Analysis / Application",
        case.get('analysis_plain', ''),
        case.get('analysis_recite', '')
    )
    
    # [C] Conclusion
    add_alac_box(doc, "C", "Conclusion", [case.get('conclusion', '')])
    add_student_explanation_box(
        doc,
        "Conclusion",
        case.get('conclusion_plain', ''),
        case.get('conclusion_recite', '')
    )
    
    # Section V
    add_heading_2(doc, "V. RELATION TO STATUTORY CONSTRUCTION")
    add_body_p(doc, case.get('syllabus', ''), bold_prefix="• Syllabus Topic Alignment:\n")
    add_latin_maxims_box(doc, case.get('latin_maxims', []))
    
    # StatCon Analysis Table
    add_statcon_table(
        doc,
        case.get('syllabus', ''),
        case.get('primary_maxim', ''),
        case.get('secondary_maxims', ''),
        case.get('violation_text', ''),
        case.get('execution_text', '')
    )
    
    add_student_explanation_box(
        doc,
        "Relation to Statutory Construction",
        case.get('statcon_plain', ''),
        case.get('statcon_recite', '')
    )
    
    # =========================================================================
    # LAYER B: STANDARDIZED 12-SECTION STATCON CASE DIGEST FORMAT
    # =========================================================================
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(8)
    p_sp.paragraph_format.space_after = Pt(4)
    
    add_template_banner(doc)
    
    p_sp2 = doc.add_paragraph()
    p_sp2.paragraph_format.space_before = Pt(4)
    p_sp2.paragraph_format.space_after = Pt(4)
    
    add_template_meta_table(doc, tmpl['title'], tmpl['gr_no'], tmpl['date'], tmpl['ponente'])
    
    # I. Facts of the Case
    add_template_section_h1(doc, "I. FACTS OF THE CASE")
    add_template_box(doc, "Material Facts of Controversy", [tmpl['facts']])
    
    # II. Facts Relevant to StatCon
    add_template_section_h1(doc, "II. FACTS RELEVANT TO STATUTORY CONSTRUCTION")
    add_template_box(doc, "Facts Giving Rise to Need for Interpretation", [
        ("• What provision of law or rule became relevant?", tmpl['facts_relevant']['provision_relevant']),
        ("• What circumstance made interpretation necessary?", tmpl['facts_relevant']['circumstance_necessary']),
        ("• What competing interpretations were presented?", tmpl['facts_relevant']['competing_interpretations'])
    ])
    
    # III. Issues
    add_template_section_h1(doc, "III. ISSUES")
    add_template_section_h2(doc, "A. General Legal Issues")
    add_body_p(doc, tmpl['issues']['general_issue'])
    add_template_section_h2(doc, "B. Statutory-Construction Issue/s")
    add_template_box(doc, "Standardized StatCon Issue Statement", [tmpl['issues']['statcon_issue']], hex_bg="EBF3FA")
    
    # IV. Provision/s of Law Construed
    add_template_section_h1(doc, "IV. PROVISION/S OF LAW CONSTRUED")
    add_template_box(doc, "Statutory / Constitutional Text & Problem", [
        ("• Quoted Provision:", tmpl['provision_construed']['quote']),
        ("• Interpretive Problem:", tmpl['provision_construed']['interpretive_problem'])
    ])
    
    # V. Ruling of the Supreme Court
    add_template_section_h1(doc, "V. RULING OF THE SUPREME COURT")
    add_template_section_h2(doc, "A. General Ruling")
    add_body_p(doc, tmpl['ruling']['general_ruling'])
    add_template_section_h2(doc, "B. Ruling on the Statutory-Construction Issue")
    add_template_box(doc, "Statutory Meaning Established by Court", [tmpl['ruling']['statcon_ruling']], hex_bg="F4FAF5")
    
    # VI. How the Court Construed the Law
    add_template_section_h1(doc, "VI. HOW THE COURT CONSTRUED THE LAW")
    add_template_box(doc, "Methodological & Doctrinal Analysis", [
        ("A. Why Was Construction Necessary?", tmpl['how_construed']['why_necessary']),
        ("B. Purpose of Construction:", tmpl['how_construed']['purpose']),
        ("C. Method, Canon, or Rule Applied:", f"{tmpl['how_construed']['canon_applied']['name']}\n   - Application: {tmpl['how_construed']['canon_applied']['application']}\n   - Result: {tmpl['how_construed']['canon_applied']['result']}"),
        ("D. Other Sources Used:", tmpl['how_construed']['other_sources'])
    ])
    
    # VII. Relevant Statutory-Construction Concepts (A to L)
    add_template_section_h1(doc, "VII. RELEVANT STATUTORY-CONSTRUCTION CONCEPTS (A–L EVALUATION)")
    add_template_concepts_table(doc, tmpl['concepts'])
    
    # VIII. Other Canons or Methods
    add_template_section_h1(doc, "VIII. OTHER CANONS OR METHODS OF CONSTRUCTION")
    add_template_box(doc, "Secondary Rules & Canons Applied", [
        ("• Canon / Maxim:", tmpl['other_canons']['name']),
        ("• How Applied:", tmpl['other_canons']['application']),
        ("• Effect on Interpretation:", tmpl['other_canons']['effect'])
    ])
    
    # IX. The Construction Made by the Court
    add_template_section_h1(doc, "IX. THE CONSTRUCTION MADE BY THE COURT")
    add_template_box(doc, "Actual Judicial Construction & Ratio", [
        ("• Construction Summary:", tmpl['construction_made']['formula']),
        ("• Principal Reasoning:", tmpl['construction_made']['reasoning'])
    ], hex_bg="F9F8F5")
    
    # X. Statutory Construction vs. Judicial Legislation
    add_template_verdict_box(doc, tmpl['statcon_or_jl']['verdict'], tmpl['statcon_or_jl']['explanation'])
    
    # XI & XII. Doctrine and One-Sentence Takeaway
    p_sp3 = doc.add_paragraph()
    p_sp3.paragraph_format.space_before = Pt(4)
    p_sp3.paragraph_format.space_after = Pt(2)
    add_template_doctrine_takeaway(doc, tmpl['doctrine'], tmpl['takeaway'])


def render_case_to_txt(case):
    """Renders a single case to plain text with both Layer A and Layer B."""
    tmpl = generate_case_digest_format_data(case)
    
    out = []
    out.append("=" * 95)
    out.append(f"CASE #{case['num']}: {case['title'].upper()}")
    out.append(f"Citation: {case['citation']}")
    out.append("=" * 95 + "\n")
    
    # LAYER A
    out.append("-----------------------------------------------------------------------------------------------")
    out.append("LAYER A: ALAC DIGEST, RECITATION SCRIPT & STATCON ANALYSIS")
    out.append("-----------------------------------------------------------------------------------------------\n")
    
    out.append("I. PARTIES & PROCEDURAL BACKGROUND")
    out.append(f"* Plaintiff/Petitioner: {case.get('plaintiff', 'N/A')}")
    out.append(f"* Defendant/Respondent: {case.get('defendant', 'N/A')}")
    out.append(f"* Originating Court: {case.get('court', 'N/A')}")
    out.append(f"* Nature of the Action: {case.get('nature', 'N/A')}\n")
    
    out.append("II. WHOLE FACTS")
    out.append(f"{case.get('facts', '')}\n")
    out.append("Plain Meaning Summary of Facts:")
    out.append(f"{case.get('plain_facts', '')}\n")
    
    out.append("III. THE LEGAL ISSUE")
    out.append(f"{case.get('issue', '')}\n")
    out.append("[HOW TO TACKLE IN RECITATION: THE LEGAL ISSUE]")
    out.append(f"* Plain Meaning: {case.get('issue_plain', '')}")
    out.append(f"* Recitation Script: {case.get('issue_recite', '')}\n")
    
    out.append("IV. ALAC DIGEST\n")
    out.append("[A] ANSWER:")
    out.append(f"{case.get('answer', '')}\n")
    out.append(f"* Plain Meaning: {case.get('answer_plain', '')}")
    out.append(f"* Recitation Script: {case.get('answer_recite', '')}\n")
    
    out.append("[L] LEGAL BASIS:")
    for lbl, val in case.get('legal_basis', []):
        out.append(f"* {lbl} {val}")
    out.append(f"* Plain Meaning: {case.get('legal_basis_plain', '')}")
    out.append(f"* Recitation Script: {case.get('legal_basis_recite', '')}\n")
    
    out.append("[A] ANALYSIS / APPLICATION:")
    for lbl, val in case.get('analysis', []):
        out.append(f"* {lbl} {val}")
    out.append(f"* Plain Meaning: {case.get('analysis_plain', '')}")
    out.append(f"* Recitation Script: {case.get('analysis_recite', '')}\n")
    
    out.append("[C] CONCLUSION:")
    out.append(f"{case.get('conclusion', '')}\n")
    out.append(f"* Plain Meaning: {case.get('conclusion_plain', '')}")
    out.append(f"* Recitation Script: {case.get('conclusion_recite', '')}\n")
    
    out.append("V. RELATION TO STATUTORY CONSTRUCTION")
    out.append(f"* Syllabus Topic Alignment:\n  {case.get('syllabus', '')}\n")
    out.append("* Governing Latin Maxims with English Translations:")
    for idx, (latin_with_eng, translation, meaning) in enumerate(case.get('latin_maxims', []), 1):
        out.append(f"  {idx}. {latin_with_eng}")
        out.append(f"     - Literal Translation: {translation}")
        out.append(f"     - StatCon Application: {meaning}")
    out.append(f"\n* Case Violation in StatCon: {case.get('violation_text', '')}")
    out.append(f"* Doctrinal Execution: {case.get('execution_text', '')}\n")
    out.append(f"* Plain StatCon Meaning: {case.get('statcon_plain', '')}")
    out.append(f"* Recitation Script: {case.get('statcon_recite', '')}\n")
    
    # LAYER B
    out.append("-----------------------------------------------------------------------------------------------")
    out.append("LAYER B: STANDARDIZED CASE DIGEST (STATUTORY CONSTRUCTION FORMAT)")
    out.append("-----------------------------------------------------------------------------------------------\n")
    out.append(f"Case Title: {tmpl['title']}")
    out.append(f"G.R. No.:   {tmpl['gr_no']}")
    out.append(f"Date:       {tmpl['date']}")
    out.append(f"Ponente:    {tmpl['ponente']}\n")
    
    out.append("I. FACTS OF THE CASE")
    out.append(f"{tmpl['facts']}\n")
    
    out.append("II. FACTS RELEVANT TO STATUTORY CONSTRUCTION")
    out.append(f"* Provision Relevant: {tmpl['facts_relevant']['provision_relevant']}")
    out.append(f"* Circumstance Making Construction Necessary: {tmpl['facts_relevant']['circumstance_necessary']}")
    out.append(f"* Competing Interpretations: {tmpl['facts_relevant']['competing_interpretations']}\n")
    
    out.append("III. ISSUES")
    out.append(f"A. General Legal Issues: {tmpl['issues']['general_issue']}")
    out.append(f"B. Statutory-Construction Issue/s: {tmpl['issues']['statcon_issue']}\n")
    
    out.append("IV. PROVISION/S OF LAW CONSTRUED")
    out.append(f"* Provision: {tmpl['provision_construed']['quote']}")
    out.append(f"* Interpretive Problem: {tmpl['provision_construed']['interpretive_problem']}\n")
    
    out.append("V. RULING OF THE SUPREME COURT")
    out.append(f"A. General Ruling: {tmpl['ruling']['general_ruling']}")
    out.append(f"B. Ruling on Statutory-Construction Issue: {tmpl['ruling']['statcon_ruling']}\n")
    
    out.append("VI. HOW THE COURT CONSTRUED THE LAW")
    out.append(f"A. Why Was Construction Necessary? {tmpl['how_construed']['why_necessary']}")
    out.append(f"B. Purpose of Construction: {tmpl['how_construed']['purpose']}")
    out.append(f"C. Method/Canon Applied: {tmpl['how_construed']['canon_applied']['name']}")
    out.append(f"   - How Applied: {tmpl['how_construed']['canon_applied']['application']}")
    out.append(f"   - Result: {tmpl['how_construed']['canon_applied']['result']}")
    out.append(f"D. Other Sources Used: {tmpl['how_construed']['other_sources']}\n")
    
    out.append("VII. RELEVANT STATUTORY-CONSTRUCTION CONCEPTS (A–L EVALUATION)")
    for c in tmpl['concepts']:
        status = "[APPLICABLE]" if c.get('applicable', True) else "[NOT APPLICABLE]"
        out.append(f"* {c['letter']}. {c['name']} - {status}: {c['explanation']}")
    out.append("")
    
    out.append("VIII. OTHER CANONS OR METHODS OF CONSTRUCTION")
    out.append(f"* Canon/Method: {tmpl['other_canons']['name']}")
    out.append(f"* How Applied: {tmpl['other_canons']['application']}")
    out.append(f"* Effect on Interpretation: {tmpl['other_canons']['effect']}\n")
    
    out.append("IX. THE CONSTRUCTION MADE BY THE COURT")
    out.append(f"* Summary Formula: {tmpl['construction_made']['formula']}")
    out.append(f"* Principal Reasoning: {tmpl['construction_made']['reasoning']}\n")
    
    out.append("X. STATUTORY CONSTRUCTION OR JUDICIAL LEGISLATION?")
    c1 = "[X]" if "proper" in tmpl['statcon_or_jl']['verdict'].lower() else "[ ]"
    c2 = "[X]" if "judicial legislation" in tmpl['statcon_or_jl']['verdict'].lower() and "proper" not in tmpl['statcon_or_jl']['verdict'].lower() else "[ ]"
    c3 = "[X]" if "debatable" in tmpl['statcon_or_jl']['verdict'].lower() else "[ ]"
    out.append(f"{c1} Proper statutory construction    {c2} Judicial legislation    {c3} Debatable")
    out.append(f"Explanation: {tmpl['statcon_or_jl']['explanation']}\n")
    
    out.append(f"XI. DOCTRINE:\n{tmpl['doctrine']}\n")
    out.append(f"XII. ONE-SENTENCE TAKEAWAY:\n{tmpl['takeaway']}\n")
    out.append("#" * 95 + "\n\n")
    
    return "\n".join(out)

print("Unified Case Renderer module ready.")
